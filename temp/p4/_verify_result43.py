# -*- coding: utf-8 -*-
"""独立复核 result4-3.xlsx: 结构/取值/恒等式 (对齐 _verify_result2.py 风格)"""
import openpyxl
import numpy as np
import pandas as pd

F = r"D:\CUMCM2026Problems\C题\p4\result4-3.xlsx"
wb = openpyxl.load_workbook(F)
print("sheets:", wb.sheetnames)

# 附件4 实际价 (slot空间) + 模板列序对齐
a4 = pd.read_excel(r"D:\CUMCM2026Problems\C题\附件\附件4.xlsx", index_col=0)
a4.columns = range(144)
lam4 = a4.values[31:365]
order = list(range(1, 144)) + [0]                     # 模板列j ↔ slot order[j]

# ---- 计划购电量 / 调整购电量 ----
for name in ("计划购电量", "调整购电量"):
    ws = wb[name]
    assert ws.max_row == 335 and ws.max_column == 147, (name, ws.max_row, ws.max_column)
    bad = neg = 0
    worst_fee = 0.0
    for d, r in enumerate(range(2, 336)):
        vals = np.array([ws.cell(row=r, column=c).value for c in range(2, 146)])
        if any(v is None for v in vals):
            bad += 1
            continue
        lam_sheet = np.array([lam4[d][order[j]] for j in range(144)])
        tot = ws.cell(row=r, column=146).value
        cost = ws.cell(row=r, column=147).value
        assert abs(round(sum(vals), 2) - tot) < 0.02, f"{name} row {r}"
        worst_fee = max(worst_fee, abs(np.sum(lam_sheet * vals) - cost))
        if min(vals) < 0:
            neg += 1
    print(f"{name}: 334行填满={bad==0}, 负值行={neg}, 全天购电费 vs Σ(对齐λ×表内值) 最大差 {worst_fee:.4f} 元")

# 调整购电量: t<36 段 q==p (slot 0..35 → 模板列: col145=slot0, col2..36=slot1..35)
ws_p, ws_q = wb["计划购电量"], wb["调整购电量"]
worst_qeqp = 0.0
cols_qeqp = [145] + list(range(2, 37))
for r in range(2, 336):
    for c in cols_qeqp:
        worst_qeqp = max(worst_qeqp, abs(ws_p.cell(row=r, column=c).value - ws_q.cell(row=r, column=c).value))
print(f"调整=计划 on 0:00-6:00段 (slot0..35): 最大差 {worst_qeqp:.6f} kWh (应≈0)")

# ---- 充放电量 ----
ws = wb["充放电量"]
print(f"充放电量: {ws.max_row} 行 (应为 1+2004=2005)")
worst = 0.0
mn = 1e18
for d in range(334):
    base = 2 + d * 6
    e0 = ws.cell(row=base, column=6).value
    e24 = ws.cell(row=base + 1, column=6).value
    chg = sum(ws.cell(row=base + s, column=3).value for s in range(6))
    dis = sum(ws.cell(row=base + s, column=4).value for s in range(6))
    worst = max(worst, abs(0.9 * chg - dis / 0.9 - (e24 - e0)))
    mn = min(mn, e0, e24)
    for s in range(6):
        assert ws.cell(row=base + s, column=3).value >= 0 and ws.cell(row=base + s, column=4).value >= 0
    assert ws.cell(row=base, column=5).value is not None and ws.cell(row=base + 1, column=5).value == "24:00"
print(f"  恒等式 0.9*Σ充-Σ放/0.9=ΔSOC 最大偏差 {worst:.4f} kWh; SOC 最低 {mn:.0f} kWh (>=1200)")
# E0链
soc0 = [ws.cell(row=2 + d * 6, column=6).value for d in range(334)]
soc24 = [ws.cell(row=2 + d * 6 + 1, column=6).value for d in range(334)]
print(f"  E0链衔接(E24(d)=E0(d+1)) 最大差: {max(abs(soc24[i]-soc0[i+1]) for i in range(333)):.4f} | 首日E0: {soc0[0]}")

# ---- 紧急购电量 ----
ws = wb["紧急购电量"]
tot = sum(ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1))
ok = all(isinstance(ws.cell(row=r, column=2).value, str) and "-" in ws.cell(row=r, column=2).value
         for r in range(2, ws.max_row + 1))
print(f"紧急购电量: {ws.max_row - 1} 行, 标签合法={ok}, 合计 {tot:,.2f} kWh")
