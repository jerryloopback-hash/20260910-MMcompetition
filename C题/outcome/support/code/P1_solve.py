# -*- coding: utf-8 -*-
"""
问题1（P1）：微网计划购电策略 —— 线性规划求解
模型与对偶分析见同目录 p1_model.md（第 9 节为 KKT/影子价格分析，供论文使用）
求解器：scipy.optimize.linprog(method='highs')，内核为 HiGHS
（论文表述：采用 HiGHS 求解器求解线性规划；环境：pip install numpy scipy openpyxl）

v2（日历对齐修订版）相对初版（v1）的改动：
1. 时段重排为日历 144 段 I_1..I_144 = [0:00,0:10)…[23:50,24:00)。附件1 行标签为时段起点
   （0:10,…,23:50,24:00），故 I_1=[0:00,0:10) 取第 144 行（0:00+1，即次日首 10 分钟）数据——
   由"每天电价/负载相同、光伏预测为同一日曲线"的逐日周期性，其数值与当日 [0:00,0:10) 相同；
   I_k（k=2..144）取第 k-1 行。
2. E(0:00)=E(24:00)=6000 双端钉死。v1 只钉 E143、E144 自由，导致跨日段"末端偷放"
   （放电 574 kWh、购电 0、E144=5362），方案不可逐日复制，且表2 回推 dSOC=-638 kWh 与
   储电量 6000/6000 矛盾。v2 消除该末端效应。
3. 模板"计划购电量"行 r（2..144）↔ I_r；行 145（0:00+1-0:10+1）填 I_1 的周期镜像值。
   表2 四小时块与日历窗口精确对齐。
4. 新增对偶/影子价格提取、KKT 阈值互补松弛校验，并用有限差分重解抽样验证对偶符号
   （p1_model.md 第 9 节）。
5. 路径改为相对本文件定位，去除硬编码盘符。
本目录 result1.xlsx / _traj.npy 已由本脚本（v2）生成；配套脚本：P1_sensitivity.py（§10.1）、
P1_greedy.py（§10.2）、P1_figs.py（§9.5 论文图）。
"""
from pathlib import Path

import numpy as np
import openpyxl
from scipy.optimize import linprog

OUTDIR = Path(__file__).resolve().parent            # …/C题/p1
BASE = OUTDIR.parent                                # …/C题

# ---------- 1. 读取附件1 ----------
wb = openpyxl.load_workbook(BASE / "附件" / "附件1.xlsx", data_only=True)
ws = wb["Sheet1"]
rows = list(ws.iter_rows(min_row=2, values_only=True))      # 144 x (时间, 电价, 负载, 光伏)
assert len(rows) == 144
raw_price = np.array([r[1] for r in rows], dtype=float)     # 元/kWh，行标签=时段起点 0:10..24:00
raw_load = np.array([r[2] for r in rows], dtype=float)      # kW
raw_pv = np.array([r[3] for r in rows], dtype=float)        # kW（预测）

DT = 1 / 6.0            # h，时段长
ETA = 0.9               # 充、放各 90%（往返 0.81）
T = 144
EMIN, EMAX = 1200.0, 10800.0
E0 = E_END = 6000.0     # 0:00 初值（附录1）；0:00 与 24:00 储电量相同 → 双端均 6000
PMAX = 5000.0           # 充/放功率上限，并网侧（母线端）口径，见模型文档注记 5

# 日历区间 I_1..I_144（数组下标 i ↔ I_{i+1} = [10i, 10(i+1)) 分钟）：
# I_1 ← 第 144 行（标签 24:00 = 次日 0:00-0:10，周期性下等于当日 0:00-0:10）；I_k ← 第 k-1 行
price = np.concatenate([raw_price[143:], raw_price[:143]])
load = np.concatenate([raw_load[143:], raw_load[:143]])
pv = np.concatenate([raw_pv[143:], raw_pv[:143]])


# ---------- 2. LP 构造 ----------
# 变量（每段 6 个）：a 购电直供, b 购电充电, g 光伏直供, c 光伏充电, d 放电, q 弃光
NV = T * 6
idx = {k: np.arange(j, NV, 6) for j, k in enumerate("abgcdq")}


def soc_row(upto):
    """前 upto+1 段的净充电累计（kWh，相对 E0）的系数行"""
    r = np.zeros(NV)
    for k in range(upto + 1):
        r[idx["b"][k]] += ETA * DT
        r[idx["c"][k]] += ETA * DT
        r[idx["d"][k]] += -DT / ETA
    return r


def build_lp(node_pin=None):
    """node_pin=(i, delta)：追加等式约束"第 i 段末净充电累计 = delta"（有限差分用）"""
    c_obj = np.zeros(NV)
    c_obj[idx["a"]] = price * DT
    c_obj[idx["b"]] = price * DT

    A_eq, b_eq = [], []
    for t in range(T):
        rL = np.zeros(NV); rL[idx["a"][t]] = 1; rL[idx["g"][t]] = 1; rL[idx["d"][t]] = 1
        rP = np.zeros(NV); rP[idx["g"][t]] = 1; rP[idx["c"][t]] = 1; rP[idx["q"][t]] = 1
        A_eq += [rL, rP]; b_eq += [load[t], pv[t]]
    A_eq.append(soc_row(T - 1)); b_eq.append(0.0)   # 全天净充放累计=0 ⇔ E(24:00)=E(0:00)=6000
    if node_pin is not None:
        i, delta = node_pin
        A_eq.append(soc_row(i)); b_eq.append(delta)

    A_ub, b_ub = [], []
    for t in range(T):                              # SOC 累计界限（所有节点 1200..10800）
        A_ub += [soc_row(t), -soc_row(t)]
        b_ub += [EMAX - E0, E0 - EMIN]
    for t in range(T):                              # 功率上限（并网侧口径）
        rC = np.zeros(NV); rC[idx["b"][t]] = 1; rC[idx["c"][t]] = 1
        rD = np.zeros(NV); rD[idx["d"][t]] = 1
        A_ub += [rC, rD]; b_ub += [PMAX, PMAX]
    return c_obj, np.array(A_eq), np.array(b_eq), np.array(A_ub), np.array(b_ub)


c_obj, A_eq, b_eq, A_ub, b_ub = build_lp()
res = linprog(c_obj, A_ub=A_ub, b_ub=b_ub, A_eq=A_eq, b_eq=b_eq,
              bounds=[(0, None)] * NV, method="highs")
assert res.success, f"求解失败: {res.message}"
x = res.x
a, b, g, c, d, q = (x[idx[k]] for k in "abgcdq")
purchase_kwh = (a + b) * DT                                  # 各段购电量 kWh

SOC = np.zeros(T + 1); SOC[0] = E0
for t in range(T):
    SOC[t + 1] = SOC[t] + ETA * (b[t] + c[t]) * DT - d[t] * DT / ETA

cost = float(res.fun)


def hm(i):
    """区间下标 i（0 基）的起点标签 hh:mm"""
    m = 10 * i
    return f"{m // 60}:{m % 60:02d}"


# ---------- 3. 结果校验 ----------
print("=== P1 v2（日历对齐）求解 ===")
print(f"最优购电费 = {cost:,.2f} 元（I_1..I_144 全天）")
print(f"全天购电量 = {purchase_kwh.sum():,.1f} kWh")
print(f"弃光总量   = {q.sum() * DT:,.2f} kWh")
print(f"同段边充边放时段数 = {(np.minimum(b + c, d) > 1e-6).sum()}")
print(f"SOC 范围 [{SOC.min():.1f}, {SOC.max():.1f}]，E(0:00)={SOC[0]:.1f}，E(24:00)={SOC[144]:.3f}")
assert abs(SOC[144] - E_END) < 1e-6, "E(24:00) != 6000"

Eload = load.sum() * DT; Epv = pv.sum() * DT; Eq = q.sum() * DT
Lch = (1 - ETA) * (b + c).sum() * DT
Ldi = (1 / ETA - 1) * d.sum() * DT
Ebuy = purchase_kwh.sum()
print(f"\n[守恒校验 全天] 购电{Ebuy:,.2f}+光伏{Epv:,.2f} = 负载{Eload:,.0f}+弃光{Eq:.2f}"
      f"+充损{Lch:,.1f}+放损{Ldi:,.1f} → 差 {Ebuy + Epv - (Eload + Eq + Lch + Ldi):.2e}")

print("\n表1 指定时段购电量（kWh）：")
for h in (600, 720, 840, 960, 1080, 1200):                  # 10,12,14,16,18,20 点
    i = h // 10
    print(f"  {hm(i)}-{hm(i + 1) if (i + 1) % 144 else '0:00+1'}: {purchase_kwh[i]:.2f}")

print("\n表2 四小时块（充电量=并网侧输入, 放电量=并网侧送达, 口径乙）：")
blocks = []
for k in range(6):
    sl = slice(24 * k, 24 * (k + 1))
    ch = (b[sl] + c[sl]).sum() * DT
    dis = d[sl].sum() * DT
    blocks.append((ch, dis))
    print(f"  {4 * k}:00-{4 * (k + 1)}:00  充 {ch:,.2f}  放 {dis:,.2f}")
print(f"  0:00/24:00 储电量 = {SOC[0]:.0f} / {SOC[144]:.0f}")
d_soc = 0.9 * sum(ch for ch, _ in blocks) - sum(dis for _, dis in blocks) / 0.9
print(f"  [表2 自洽校验] 0.9*Σ充 - Σ放/0.9 = {d_soc:,.2f} kWh（应≈0 = E(24:00)-E(0:00)）")

# ---------- 4. 对偶 / 影子价格（模型文档第 9 节） ----------
have_dual = res.ineqlin is not None and res.ineqlin.marginals is not None
if have_dual:
    m_ineq = res.ineqlin.marginals                          # ∂f/∂b_ub ≤ 0
    mU = m_ineq[0:2 * T:2]                                  # SOC 上界行
    mL = m_ineq[1:2 * T:2]                                  # SOC 下界行
    mC = m_ineq[2 * T::2]                                   # 充电功率行
    mD = m_ineq[2 * T + 1::2]                               # 放电功率行
    mPin = res.eqlin.marginals[-1]                          # 全天归零约束（最后一条等式）
    # 节点边际价值 γ_k（= 节点 k 末白得 1 kWh 可省费用，≥0）：
    #   递推 γ_k = γ_{k+1} - η_k + θ_k（内点不变；满罐 η>0 下台阶；空罐 θ>0 上台阶），终端 γ_{T+1}=mPin
    #   累计形式：γ_k = mPin + Σ_{t≥k}(mU_t - mL_t)   [mU=-η ≤0, mL=-θ ≤0]
    sfx = np.cumsum((mU - mL)[::-1])[::-1]
    psi = mPin + sfx                                        # 元/kWh，下标 k=0..143 ↔ 节点 1..144

    # KKT 互补松弛校验（σ_k = 负载行对偶/Δt = 边际供电成本；a_k>0 时 σ_k=λ_k）：
    #   充电段: λ_k = 0.9γ_k - π_k/Δt → 0.9γ_k - λ_k = π_k/Δt ≥ 0（顶格时>0，否则=0）
    #   放电段: σ_k = γ_k/0.9 + ρ_k/Δt → σ_k - γ_k/0.9 = ρ_k/Δt ≥ 0（顶格时>0，否则=0）
    #   空闲段: 迟滞带 0.9γ_k - π_k/Δt ≤ λ_k 且 σ_k ≤ γ_k/0.9 + ρ_k/Δt
    ch_on = (b + c) > 1e-6; di_on = d > 1e-6
    pi, rho = -mC / DT, -mD / DT
    sigma = res.eqlin.marginals[0:2 * T:2] / DT             # 负载行在偶数位（0,2,4,...）
    viol_ch = np.where(ch_on, 0.9 * psi - price - pi, np.minimum(0.9 * psi - price - pi, 0.0))
    viol_di = np.where(di_on, sigma - psi / 0.9 - rho, np.minimum(sigma - psi / 0.9 - rho, 0.0))
    print(f"\n[对偶校验] 互补松弛最大违约: 充电侧 {viol_ch.max():.2e}, 放电侧 {viol_di.max():.2e}（元/kWh）")
    print(f"Ψ̂ 范围 [{psi.min():.3f}, {psi.max():.3f}] 元/kWh；"
          f"充电功率顶格段 {int((b + c > PMAX - 1e-6).sum())} 个（其影子溢价 π>0）")

    # 有限差分交叉验证（抽样内点节点）：在节点 i 挂一个零成本、上限 δ 的"自由能量"变量，
    # 目标函数斜率应 ≈ -γ_i（白得 1 kWh 的价值）。注意不能用"钉住 P_i=δ"——那是强迫多充，
    # 含获取成本，不是白得能量。
    def solve_with_free_energy(i, delta):
        """追加变量 f∈[0,δ]（零成本），使 P_t += f 对所有 t≥i 成立"""
        c2 = np.concatenate([c_obj, [0.0]])
        col_eq = np.zeros(A_eq.shape[0])
        col_eq[-1] = 1.0                                   # 全天归零行（=soc_row(T-1)）
        Ae = np.column_stack([A_eq, col_eq]); be = b_eq.copy()
        col_ub = np.zeros(A_ub.shape[0])
        for t in range(i, T):                              # 前缀行 t≥i：偶=上界行，奇=下界行
            col_ub[2 * t] += 1.0; col_ub[2 * t + 1] -= 1.0
        Au = np.column_stack([A_ub, col_ub]); bu = b_ub.copy()
        return linprog(c2, A_ub=Au, b_ub=bu, A_eq=Ae, b_eq=be,
                       bounds=[(0, None)] * NV + [(0, delta)], method="highs")

    interior = [i for i in range(T) if EMIN + 10 < SOC[i + 1] < EMAX - 10]
    picks = [interior[j] for j in np.linspace(0, len(interior) - 1, 8).astype(int)]
    devs = []
    for i in picks:
        r2 = solve_with_free_energy(i, 5.0)
        if r2.success:
            devs.append(abs(psi[i] + (r2.fun - cost) / 5.0))
    print(f"[对偶符号有限差分验证] 抽样 {len(devs)} 个内点节点，最大偏差 {max(devs):.2e} 元/kWh")

    # 充/放窗口表（论文第 9 节用）
    def windows(mask):
        out, s = [], None
        for i in range(T):
            if mask[i] and s is None:
                s = i
            elif not mask[i] and s is not None:
                out.append((s, i - 1)); s = None
        if s is not None:
            out.append((s, T - 1))
        return out

    def endlabel(e):
        return hm(e + 1) if (e + 1) % 144 else "24:00"

    print("\n放电窗口（起止, 电量kWh, λ范围, Ψ̂范围）:")
    for s, e in windows(di_on):
        print(f"  {hm(s)}-{endlabel(e)}: {d[s:e+1].sum()*DT:8.1f}"
              f"  λ [{price[s:e+1].min():.3f},{price[s:e+1].max():.3f}]"
              f"  Ψ̂ [{psi[s:e+1].min():.3f},{psi[s:e+1].max():.3f}]")
    print("净充电窗口（起止, 电量kWh, λ范围, Ψ̂范围）:")
    for s, e in windows(ch_on):
        print(f"  {hm(s)}-{endlabel(e)}: {(b+c)[s:e+1].sum()*DT:8.1f}"
              f"  λ [{price[s:e+1].min():.3f},{price[s:e+1].max():.3f}]"
              f"  Ψ̂ [{psi[s:e+1].min():.3f},{psi[s:e+1].max():.3f}]")
else:
    psi = None
    print("\n[对偶] 当前 scipy 未返回 marginals，跳过对偶分析（需 scipy ≥ 1.7 的 highs 方法）")

# ---------- 5. 写入 result1.xlsx ----------
tpl = openpyxl.load_workbook(BASE / "附件" / "附件5" / "result1.xlsx")
ws1, ws2 = tpl["计划购电量"], tpl["充放电量"]
for r in range(2, 146):                                     # 行 r ↔ I_r；行145 = I_1 周期镜像
    i = r - 1 if r <= 144 else 0
    ws1.cell(row=r, column=2).value = round(float(purchase_kwh[i]), 2)
for k in range(6):
    ch, dis = blocks[k]
    ws2.cell(row=2 + k, column=2).value = round(ch, 2)
    ws2.cell(row=2 + k, column=3).value = round(dis, 2)
ws2.cell(row=2, column=5).value = round(SOC[0], 2)          # 0:00 储电量
ws2.cell(row=3, column=5).value = round(SOC[144], 2)        # 24:00 储电量
out = OUTDIR / "result1.xlsx"
tpl.save(out)
print(f"\n已保存 -> {out}")

# 舍入后复检（以发布数值回推）
pub = np.array([round(float(v), 2) for v in purchase_kwh])
print(f"[舍入复检] 发布值全天购电 {pub.sum():,.2f} kWh，与未舍入差 {pub.sum() - Ebuy:+.2e} kWh")

# 存档轨迹（行序：a,b,g,c,d,q,a+b,SOC[1:],psi）
traj = [a, b, g, c, d, q, a + b, SOC[1:]] + ([psi] if psi is not None else [])
np.save(OUTDIR / "_traj.npy", np.vstack(traj))
