# -*- coding: utf-8 -*-
"""result2 模板列头与示例行完整倾倒"""
import openpyxl

wb = openpyxl.load_workbook(r"附件\附件5\result2.xlsx")

ws = wb["计划购电量"]
hdr = [c.value for c in ws[1]]
print(f"计划购电量: {len(hdr)} 列")
for i, h in enumerate(hdr):
    print(f"  col {i+1:3d} ({openpyxl.utils.get_column_letter(i+1)}): {h!r}")

for name in ["充放电量", "紧急购电量"]:
    ws = wb[name]
    print(f"\n===== {name}: {ws.max_row} 行 x {ws.max_column} 列 =====")
    for r in ws.iter_rows(values_only=True):
        print("   ", r)
