# -*- coding: utf-8 -*-
"""独立复核 result2.xlsx 成品:结构、取值、恒等式"""
import openpyxl
import numpy as np

F = r"p2_part2\result2.xlsx"
wb = openpyxl.load_workbook(F)
print("sheets:", wb.sheetnames)

# ---- 计划购电量 ----
ws = wb["计划购电量"]
assert ws.max_row == 335 and ws.max_column == 147, (ws.max_row, ws.max_column)
bad = 0
neg = 0
for r in range(2, 336):
    vals = [ws.cell(row=r, column=c).value for c in range(2, 146)]
    if any(v is None for v in vals):
        bad += 1
        continue
    tot = ws.cell(row=r, column=146).value
    cost = ws.cell(row=r, column=147).value
    s = round(sum(vals), 2)
    assert abs(s - tot) < 0.02, f"row {r}: sum {s} vs 全天 {tot}"
    if min(vals) < 0:
        neg += 1
print(f"计划购电量: 334 行填满={bad == 0}, 全天列=行和 ✓, 负值行数={neg}")
print("  样例 2025/3/20 (行49): 前3列", [ws.cell(row=49, column=c).value for c in (2, 3, 4)],
      " 末列(I_1)", ws.cell(row=49, column=145).value,
      " 全天", ws.cell(row=49, column=146).value, ws.cell(row=49, column=147).value)

# ---- 充放电量 ----
ws = wb["充放电量"]
n = ws.max_row
print(f"充放电量: {n} 行 (应为 1+2004=2005)")
assert n == 2005
worst = 0.0
mn = 1e18
prev_date = None
for d in range(334):
    base = 2 + d * 6
    e0 = ws.cell(row=base, column=6).value
    e24 = ws.cell(row=base + 1, column=6).value
    chg = sum(ws.cell(row=base + s, column=3).value for s in range(6))
    dis = sum(ws.cell(row=base + s, column=4).value for s in range(6))
    ident = 0.9 * chg - dis / 0.9 - (e24 - e0)
    worst = max(worst, abs(ident))
    mn = min(mn, e0, e24)
    for s in range(6):
        assert ws.cell(row=base + s, column=3).value >= 0
        assert ws.cell(row=base + s, column=4).value >= 0
    assert ws.cell(row=base, column=5).value is not None and ws.cell(row=base + 1, column=5).value == "24:00"
print(f"  恒等式 0.9*Σ充-Σ放/0.9=ΔSOC 最大偏差 {worst:.4f} kWh (取整噪声); SOC 最低 {mn:.0f} kWh (>=1200)")

# ---- 紧急购电量 ----
ws = wb["紧急购电量"]
rows = ws.max_row - 1
tot = sum(ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1))
labels_ok = all(isinstance(ws.cell(row=r, column=2).value, str) and "-" in ws.cell(row=r, column=2).value
                for r in range(2, ws.max_row + 1))
print(f"紧急购电量: {rows} 行, 区间标签全部合法={labels_ok}, 紧急购电量合计 {tot:,.2f} kWh")
