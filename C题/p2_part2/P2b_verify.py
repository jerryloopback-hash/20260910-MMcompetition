# -*- coding: utf-8 -*-
"""独立复核 result2.xlsx 成品(v2 口径):结构、日期覆盖、10分钟紧急购电网格、恒等式
运行目录: C题/ 根目录 (python p2_part2/P2b_verify.py)"""
import openpyxl

F = r"p2_part2\result2.xlsx"
wb = openpyxl.load_workbook(F)
print("sheets:", wb.sheetnames)


def hdr_map(ws):
    return {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}


def find_col(ws, hd, *keys, exclude=()):
    hits = [c for h, c in sorted(hd.items(), key=lambda kv: kv[1])
            if all(k in h for k in keys) and not any(e in h for e in exclude)]
    assert len(hits) == 1, f"[{ws.title}] 表头定位 {keys} -> {hits}"
    return hits[0]


def fc(m):
    return "24:00" if m >= 1440 else f"{m // 60}:{m % 60:02d}"


LABELS = [f"{fc(s * 10)}-{fc((s + 1) * 10)}" for s in range(144)]

# ---- 计划购电量 ----
ws = wb["计划购电量"]
hd = hdr_map(ws)
col_day = next((c for h, c in hd.items() if "日期" in h), None)
col_tot = find_col(ws, hd, "全天购电量")
col_cost = find_col(ws, hd, "全天购电费")
assert ws.max_row == 335, ws.max_row
bad = neg = 0
for r in range(2, 336):
    vals = [ws.cell(row=r, column=c).value for c in range(2, col_tot)]
    if any(v is None for v in vals):
        bad += 1
        continue
    s = round(sum(vals), 2)
    assert abs(s - ws.cell(row=r, column=col_tot).value) < 0.02, f"row {r}: sum {s}"
    if min(vals) < 0:
        neg += 1
date_ok = None
if col_day is not None:
    dts = [ws.cell(row=r, column=col_day).value for r in range(2, 336)]
    date_ok = (all(v is not None for v in dts)
               and dts[0].strftime("%Y-%m-%d") == "2025-02-01"
               and dts[-1].strftime("%Y-%m-%d") == "2025-12-31")
print(f"计划购电量: 334 行填满={bad == 0}, 全天列=行和 ✓, 负值行数={neg}, "
      f"日期列逐行填写={date_ok}")
print("  样例 2025/3/20 (行49): 前3列", [ws.cell(row=49, column=c).value for c in (2, 3, 4)],
      " 末列(I_1)", ws.cell(row=49, column=col_tot - 1).value,
      " 全天", ws.cell(row=49, column=col_tot).value, ws.cell(row=49, column=col_cost).value)

# ---- 充放电量 ----
ws = wb["充放电量"]
hd = hdr_map(ws)
col_day = find_col(ws, hd, "日期")
col_chg = find_col(ws, hd, "充电量", exclude=("放",))
col_dis = find_col(ws, hd, "放电量")
soc_cols = sorted(c for h, c in hd.items() if "储电量" in h)
assert len(soc_cols) >= 1, f"充放电量储电量列头缺失: {hd}"
two_col_mode = len(soc_cols) >= 2
n = ws.max_row
assert n == 2005, n
worst = 0.0
mn = 1e18
date_bad = 0
for d in range(334):
    base = 2 + d * 6
    dts = [ws.cell(row=base + s, column=col_day).value for s in range(6)]
    if any(v is None for v in dts) or len({v.date() for v in dts}) != 1:
        date_bad += 1
    if two_col_mode:
        e0 = ws.cell(row=base, column=soc_cols[0]).value
        e24 = ws.cell(row=base, column=soc_cols[1]).value
    else:
        assert ws.cell(row=base, column=soc_cols[0] - 1).value == "0:00"
        assert ws.cell(row=base + 1, column=soc_cols[0] - 1).value == "24:00"
        e0 = ws.cell(row=base, column=soc_cols[0]).value
        e24 = ws.cell(row=base + 1, column=soc_cols[0]).value
    chg = sum(ws.cell(row=base + s, column=col_chg).value for s in range(6))
    dis = sum(ws.cell(row=base + s, column=col_dis).value for s in range(6))
    worst = max(worst, abs(0.9 * chg - dis / 0.9 - (e24 - e0)))
    mn = min(mn, e0, e24)
    for s in range(6):
        assert ws.cell(row=base + s, column=col_chg).value >= 0
        assert ws.cell(row=base + s, column=col_dis).value >= 0
print(f"充放电量: {n} 行 ✓, 日期逐行覆盖且同日一致={date_bad == 0}, "
      f"恒等式 0.9*Σ充-Σ放/0.9=ΔSOC 最大偏差 {worst:.4f} kWh; SOC 最低 {mn:.0f} kWh (>=1200)")

# ---- 紧急购电量 ----
ws = wb["紧急购电量"]
hd = hdr_map(ws)
col_day = find_col(ws, hd, "日期")
col_itv = find_col(ws, hd, "时间段")
col_amt = find_col(ws, hd, "购电量", exclude=("计划", "全天"))
assert ws.max_row == 1 + 334 * 144, ws.max_row
tot = 0.0
nz = 0
first = last = None
grid_ok = True
r = 1
for d in range(334):
    dv0 = None
    for s in range(144):
        r += 1
        dv = ws.cell(row=r, column=col_day).value
        assert dv is not None, f"行{r} 日期缺失"
        if s == 0:
            dv0 = dv
            if first is None:
                first = dv
            last = dv
        elif dv.date() != dv0.date():
            grid_ok = False
        if ws.cell(row=r, column=col_itv).value != LABELS[s]:
            grid_ok = False
        v = ws.cell(row=r, column=col_amt).value
        assert v is not None and v >= 0, f"行{r} 购电量异常: {v}"
        tot += v
        nz += v > 0
print(f"紧急购电量: 48096 行 10分钟网格 ✓(每行带日期且时段标签逐行对齐={grid_ok}), "
      f"日期 {first:%Y-%m-%d} .. {last:%Y-%m-%d}, "
      f"非零 {nz} 时段, 合计 {tot:,.2f} kWh")
