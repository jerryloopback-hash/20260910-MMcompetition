# -*- coding: utf-8 -*-
"""
生成问题2交付文件 result2.xlsx (2026-09-12)
================================================
口径与 p2b_quantile_sweep.py / p2b_solve.py 完全一致:
  预测输入 = EGD集成v5点预测 + 过去30天同时点残差经验分位
  决策分位 tau = 0.81 (论文 6.2.4 全年费用回测最优): 负载取 P_tau, 光伏取 P_{1-tau}
  计划层 = p2b_solve.solve_day(LP, 含日末储能价值项), 执行结算与 p2b_solve.settle 同逻辑
  E0 滚动: E0^(d+1) = 执行层末SOC, 首日 6000 kWh
输出(模板 = 附件/附件5/result2.xlsx, 成品另存 p2_part2/result2.xlsx):
  计划购电量: 模板列序 [I_2..I_144 | I_1](kWh) + 全天购电量 + 全天购电费(按表内取整值计)
  充放电量:   每日 6 块的执行充电/放电量(kWh) + 0:00/24:00 执行储电量
  紧急购电量: 连续缺口区间合并记账(表4 格式, kWh)
自检:
  1) 逐日 settle_full 与 p2b_solve.settle 的计划费/紧急费/末SOC 逐项一致
  2) 全年 plan/emerg/total 与 p2b_分位遍历存档中 tau=0.81 行对账
  3) 储能块恒等式 0.9*充电 - 放电/0.9 = 末SOC - 初SOC
用法: python p2b_make_result2.py
"""
import time
from datetime import time as dtime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
import openpyxl

HERE = Path(__file__).resolve().parent
sys_path_added = False
if not sys_path_added:
    import sys
    sys.path.insert(0, str(HERE))
    sys_path_added = True

from p2b_solve import (solve_day, settle, lam, Lmat, Pmat, dates, i_feb1,
                       T, DT, ETA, EMIN, EMAX, PMAX, PRED)  # noqa: E402

TAU = 0.81
N = 334
ROLLING_WINDOW = 30
TEMPLATE = HERE.parent / "附件" / "附件5" / "result2.xlsx"
OUT = HERE / "result2.xlsx"
SPEC_DATES = ["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"]

# ---------- 1. 分位修正预测(与 p2b_quantile_sweep.rolling_quantile_at 同法) ----------
load_point = np.load(PRED / "load_pred_ens_v5.npy")
pv_point = np.load(PRED / "pv_pred_ens_v5.npy")
assert load_point.shape == (N, T) and pv_point.shape == (N, T)

pred_idx = np.arange(i_feb1, i_feb1 + N)
pred_set = set(pred_idx.tolist())
hrs = (np.arange(T) + 1) * 10.0 / 60.0
night = (hrs <= 6.0) | (hrs >= 20.0)


def rolling_quantile_at(point_pred, true_mat, q, is_pv):
    """点预测 + 过去30天同时点残差的经验 q 分位; 冷启动期用 persistence。"""
    out = np.zeros((N, T))
    for k, i_day in enumerate(pred_idx):
        lo = max(0, i_day - ROLLING_WINDOW)
        res = []
        for j in range(lo, i_day):
            if j in pred_set:
                pj = point_pred[j - i_feb1]
            else:
                pj = true_mat[j - 1] if j > 0 else true_mat[j]
            res.append(true_mat[j] - pj)
        out[k] = point_pred[k] + np.quantile(np.asarray(res), q, axis=0)
        if is_pv:
            out[k, night] = 0.0
        out[k] = np.clip(out[k], 0, None)
    return out


# ---------- 2. 执行结算(与 p2b_solve.settle 同逻辑, 增加逐时段输出) ----------
def settle_full(a, b, gp, dp, Lact, Pact, E0):
    E = E0
    pcost = ecost = gapE = 0.0
    chin_e = np.empty(T)
    dis_e = np.empty(T)
    gap_e = np.empty(T)
    for t in range(T):
        gr = gp[t] if gp[t] < Pact[t] else Pact[t]          # 实际光伏先满足计划直供
        da = dp[t]
        cap = ETA * (E - EMIN) / DT                         # 1200 底线截断
        if da > cap:
            da = cap
        S = a[t] + gr + da
        gap = Lact[t] - S
        gap_kw = gap if gap > 0 else 0.0
        if gap_kw > 0:
            ecost += 5 * lam[t] * gap_kw * DT
            gapE += gap_kw * DT
        cr = Pact[t] - gp[t]
        if cr < 0:
            cr = 0.0
        chin = (b[t] + cr) * DT                             # 充电=计划购电充电+实现光伏盈余
        if chin > PMAX * DT:
            chin = PMAX * DT
        cap2 = (EMAX - E) / ETA
        if chin > cap2:
            chin = cap2
        chin_e[t] = chin
        dis_e[t] = da * DT
        gap_e[t] = gap_kw * DT
        E += ETA * chin - da * DT / ETA
        pcost += lam[t] * (a[t] + b[t]) * DT                # 购电按计划量计费
    return pcost, ecost, gapE, chin_e, dis_e, gap_e, E


# ---------- 3. 全年滚动 ----------
def fmt_clock(m):
    """分钟 -> 'H:MM'(不含前导零); 1440 -> '24:00'"""
    return "24:00" if m >= 1440 else f"{m // 60}:{m % 60:02d}"


if __name__ == "__main__":
    t0 = time.time()
    print(f"[1/3] 构造 tau={TAU} 分位修正预测 ...")
    qL = rolling_quantile_at(load_point, Lmat, TAU, False)
    qP = rolling_quantile_at(pv_point, Pmat, 1.0 - TAU, True)

    print("[2/3] 逐日 LP + 执行结算 ...")
    plan_rows = []          # 每日 144 值, 模板列序 [I_2..I_144 | I_1]
    block_rows = []         # 每日 6 块 (chg, dis) + (E0, E24)
    gap_rows = []           # (date, 'H:MM-H:MM', kWh)
    spec_pick = {}          # 指定日期 -> (全天购电量, 全天购电费, blocks, E0, E24, gaps)
    E0 = 6000.0
    tot_plan = tot_emerg = 0.0
    for k in range(N):
        drow = i_feb1 + k
        Lact = np.concatenate([Lmat[drow - 1, 143:144], Lmat[drow, 0:143]])
        Pact = np.concatenate([Pmat[drow - 1, 143:144], Pmat[drow, 0:143]])
        if k == 0:
            fcl_prev = Lmat[drow - 1, 142:143]
            fcp_prev = Pmat[drow - 1, 142:143]
        else:
            fcl_prev = qL[k - 1, 143:144]
            fcp_prev = qP[k - 1, 143:144]
        Lf = np.maximum(np.concatenate([fcl_prev, qL[k, 0:143]]), 0.0)
        Pf = np.maximum(np.concatenate([fcp_prev, qP[k, 0:143]]), 0.0)

        a, b, gp, cp, dp = solve_day(Lf, Pf, E0)
        pc, ec, ge, chin_e, dis_e, gap_e, Eend = settle_full(a, b, gp, dp, Lact, Pact, E0)
        # 自检1: 与 p2b_solve.settle 逐项一致
        p1, e1, g1, n1, sn, sx, Eend2 = settle(a, b, gp, dp, Lact, Pact, E0)
        assert abs(pc - p1) < 1e-6 and abs(ec - e1) < 1e-6 and abs(Eend - Eend2) < 1e-6, f"day {k} 结算不一致"
        assert abs(ge - g1) < 1e-6 and gap_e.sum() - g1 < 1e-6

        plan_e = (a + b) * DT                                   # kWh, I_1..I_144
        vals = np.concatenate([plan_e[1:], plan_e[:1]])         # 模板列序 [I_2..I_144 | I_1]
        vals_r = np.round(vals, 2) + 0.0                        # +0.0 归一化 -0.0(LP负零噪声)
        lam_sheet = np.concatenate([lam[1:], lam[:1]])
        day_cost = float(np.sum(lam_sheet * vals_r))            # 全天购电费按表内取整值计
        day_energy = float(vals_r.sum())
        plan_rows.append((vals_r, round(day_energy, 2), round(day_cost, 2)))

        chg_blk = [round(float(chin_e[s * 24:(s + 1) * 24].sum()), 2) + 0.0 for s in range(6)]
        dis_blk = [round(float(dis_e[s * 24:(s + 1) * 24].sum()), 2) + 0.0 for s in range(6)]
        # 自检3: 储能块恒等式(未取整口径)
        lhs = ETA * chin_e.sum() - dis_e.sum() / ETA
        assert abs(lhs - (Eend - E0)) < 1e-6, f"day {k} SOC 恒等式不成立"
        block_rows.append((chg_blk, dis_blk, round(E0, 2), round(Eend, 2)))

        d = dates[drow]
        idx = np.where(gap_e > 1e-9)[0]
        day_gaps = []
        if idx.size:
            brk = np.where(np.diff(idx) > 1)[0]
            for s0, se in zip(np.concatenate([[0], brk + 1]), np.concatenate([brk, [idx.size - 1]])):
                s, e = idx[s0], idx[se]
                label = f"{fmt_clock(s * 10)}-{fmt_clock((e + 1) * 10)}"
                day_gaps.append((label, round(float(gap_e[s:e + 1].sum()), 2)))
            gap_rows.append((d, day_gaps))

        date_str = pd.Timestamp(d).strftime("%Y-%m-%d")
        if date_str in SPEC_DATES:
            spec_pick[date_str] = (day_energy, round(day_cost, 2), chg_blk, dis_blk,
                                   round(E0, 2), round(Eend, 2), day_gaps)
        tot_plan += pc
        tot_emerg += ec
        E0 = Eend
        if (k + 1) % 120 == 0:
            print(f"    {k + 1}/{N} 天  累计 计划{tot_plan:,.0f} 紧急{tot_emerg:,.0f} 元  ({time.time() - t0:.0f}s)")

    terminal_value = ETA * lam[-1] * E0
    net = tot_plan + tot_emerg - terminal_value
    print(f"[对账] 全年 计划 {tot_plan:,.2f} + 紧急 {tot_emerg:,.2f} - 储能价值 {terminal_value:,.2f}"
          f" = 净成本 {net:,.2f} 元 ({net / 1e4:.2f} 万)")

    sweep_file = HERE / "p2b_分位遍历_0.78_0.86.xlsx"
    ref = pd.read_excel(sweep_file)
    ref_row = ref[np.isclose(ref["tau"].astype(float), TAU)]
    assert len(ref_row) == 1, f"遍历存档中未找到 tau={TAU}"
    rr = ref_row.iloc[0]
    print(f"[对账] 遍历存档 tau={TAU}: 计划 {rr['plan']:,.2f} 紧急 {rr['emerg']:,.2f} 净 {rr['total']:,.2f}")
    assert abs(tot_plan - rr["plan"]) < 0.5 and abs(tot_emerg - rr["emerg"]) < 0.5 \
        and abs(net - rr["total"]) < 0.5, "与遍历存档对账失败"
    print("[对账] 与遍历存档一致 ✓")

    # ---------- 4. 写 result2.xlsx ----------
    print("[3/3] 写入 result2.xlsx ...")
    wb = openpyxl.load_workbook(TEMPLATE)

    ws = wb["计划购电量"]
    for k in range(N):
        r = k + 2
        vals_r, day_energy, day_cost = plan_rows[k]
        for j, v in enumerate(vals_r):
            ws.cell(row=r, column=2 + j, value=float(v))
        ws.cell(row=r, column=146, value=day_energy)
        ws.cell(row=r, column=147, value=day_cost)

    ws = wb["充放电量"]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for k in range(N):
        d = dates[i_feb1 + k]
        chg_blk, dis_blk, e0, e24 = block_rows[k]
        base = ws.max_row + 1
        for s in range(6):
            r = base + s
            if s == 0:
                ws.cell(row=r, column=1, value=pd.Timestamp(d).to_pydatetime())
            ws.cell(row=r, column=2, value=["0:00-4:00", "4:00-8:00", "8:00-12:00",
                                            "12:00-16:00", "16:00-20:00", "20:00-24:00"][s])
            ws.cell(row=r, column=3, value=chg_blk[s])
            ws.cell(row=r, column=4, value=dis_blk[s])
            if s == 0:
                ws.cell(row=r, column=5, value=dtime(0, 0))
                ws.cell(row=r, column=6, value=e0)
            elif s == 1:
                ws.cell(row=r, column=5, value="24:00")
                ws.cell(row=r, column=6, value=e24)

    ws = wb["紧急购电量"]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for d, day_gaps in gap_rows:
        base = ws.max_row + 1
        for i, (label, kwh) in enumerate(day_gaps):
            r = base + i
            if i == 0:
                ws.cell(row=r, column=1, value=pd.Timestamp(d).to_pydatetime())
            ws.cell(row=r, column=2, value=label)
            ws.cell(row=r, column=3, value=kwh)

    wb.save(OUT)
    n_gap_days = len(gap_rows)
    n_gap_itv = sum(len(g) for _, g in gap_rows)
    print(f"\n已保存 -> {OUT}")
    print(f"紧急购电: {n_gap_days} 天有缺口, 共 {n_gap_itv} 个连续区间; "
          f"全年购电量合计 {sum(r[1] for r in plan_rows):,.2f} kWh")
    for ds in SPEC_DATES:
        de, dc, cb, db, e0, e24, gs = spec_pick[ds]
        print(f"\n指定日期 {ds}: 全天购电 {de:,.2f} kWh, 计划费 {dc:,.2f} 元, "
              f"0:00 SOC {e0:,.0f} -> 24:00 SOC {e24:,.0f} kWh")
        print(f"  六块充电 {cb}")
        print(f"  六块放电 {db}")
        print(f"  紧急购电 {len(gs)} 段: {gs if len(gs) <= 8 else gs[:8] + [('...', '...')]}")
    print(f"\n完成, 耗时 {time.time() - t0:.0f}s")
